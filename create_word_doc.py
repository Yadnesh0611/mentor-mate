import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_hyperlink(paragraph, url, text, color="0284C7", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

doc = docx.Document()

# Document Title
title_p = doc.add_paragraph()
r_title = title_p.add_run("Mentor Mate — Research References & Official Citations")
r_title.bold = True
r_title.font.size = Pt(20)
r_title.font.color.rgb = RGBColor(15, 23, 42)
r_title.font.name = 'Segoe UI'

sub_p = doc.add_paragraph()
r_sub = sub_p.add_run("Smart India Hackathon (SIH 2026) | Team Mentor Mate | PS ID: SIH26207\n")
r_sub.font.size = Pt(11)
r_sub.font.color.rgb = RGBColor(100, 116, 139)
r_sub.font.name = 'Segoe UI'

# Section 1
h1 = doc.add_heading(level=1)
r1 = h1.add_run("1. Active Recall & Practice Testing")
r1.font.name = 'Segoe UI'
r1.font.color.rgb = RGBColor(2, 132, 199)

p1 = doc.add_paragraph()
p1.add_run("• Research Paper: ").bold = True
p1.add_run("Improving Students' Learning with Effective Learning Techniques: Promising Directions From Cognitive and Educational Psychology\n")
p1.add_run("• Authors: ").bold = True
p1.add_run("John Dunlosky, Katherine A. Rawson, Elizabeth J. Marsh, Mitchell J. Nathan, Daniel T. Willingham (2013)\n")
p1.add_run("• Journal: ").bold = True
p1.add_run("Psychological Science in the Public Interest, 14(1), 4–58\n")
p1.add_run("• Official DOI Link: ").bold = True
add_hyperlink(p1, "https://doi.org/10.1177/1529100612453266", "https://doi.org/10.1177/1529100612453266")
p1.add_run("\n• Open Access Paper: ").bold = True
add_hyperlink(p1, "https://www.psychologicalscience.org/publications/journals/pspi/learning-techniques.html", "Association for Psychological Science Article")

# Section 2
h2 = doc.add_heading(level=1)
r2 = h2.add_run("2. Guided Socratic Tutoring & Stepped Hints")
r2.font.name = 'Segoe UI'
r2.font.color.rgb = RGBColor(99, 102, 241)

p2 = doc.add_paragraph()
p2.add_run("• Paper 1 (Socratic Tutoring): ").bold = True
p2.add_run("Learning from Human Tutoring\n")
p2.add_run("  - Authors: ").bold = True
p2.add_run("Michelene T.H. Chi, Stephanie A. Siler, Heisawn Jeong, Takashi Yamauchi, Robert G.M. Hausmann (2001)\n")
p2.add_run("  - Journal: ").bold = True
p2.add_run("Cognitive Science, 25(4), 471–533\n")
p2.add_run("  - Official DOI Link: ").bold = True
add_hyperlink(p2, "https://doi.org/10.1207/s15516709cog2504_1", "https://doi.org/10.1207/s15516709cog2504_1")

p2_b = doc.add_paragraph()
p2_b.add_run("• Paper 2 (1-on-1 Tutoring Effectiveness): ").bold = True
p2_b.add_run("The 2 Sigma Problem: The Search for Methods of Group Instruction as Effective as One-to-One Tutoring\n")
p2_b.add_run("  - Author: ").bold = True
p2_b.add_run("Benjamin S. Bloom (1984)\n")
p2_b.add_run("  - Journal: ").bold = True
p2_b.add_run("Educational Researcher, 13(6), 4–16\n")
p2_b.add_run("  - Official Link: ").bold = True
add_hyperlink(p2_b, "https://doi.org/10.2307/1175554", "https://doi.org/10.2307/1175554")

# Section 3
h3 = doc.add_heading(level=1)
r3 = h3.add_run("3. National Curriculum Frameworks & Open Repositories")
r3.font.name = 'Segoe UI'
r3.font.color.rgb = RGBColor(16, 185, 129)

p3 = doc.add_paragraph()
p3.add_run("• NCERT Textbooks Portal: ").bold = True
add_hyperlink(p3, "https://ncert.nic.in/textbook.php", "https://ncert.nic.in/textbook.php")
p3.add_run("\n• National Curriculum Framework (NCF): ").bold = True
add_hyperlink(p3, "https://ncert.nic.in/national-curriculum-framework.php", "https://ncert.nic.in/national-curriculum-framework.php")
p3.add_run("\n• NPTEL Courses (IITs & IISc): ").bold = True
add_hyperlink(p3, "https://nptel.ac.in/", "https://nptel.ac.in/")
p3.add_run("\n• SWAYAM Central Portal: ").bold = True
add_hyperlink(p3, "https://swayam.gov.in/", "https://swayam.gov.in/")

output_file = 'Mentor_Mate_Research_References.docx'
doc.save(output_file)
print("Successfully generated Word document:", os.path.abspath(output_file))
